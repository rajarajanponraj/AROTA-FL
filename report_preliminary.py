import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx_helpers import add_p, add_heading_1, add_heading_2, set_cell_background, set_cell_margins, set_table_borders

def build_preliminary_pages(doc):
    # -------------------------------------------------------------
    # 1. COVER / TITLE PAGE
    # -------------------------------------------------------------
    p_title_top = add_p(doc, "A PROJECT PHASE REPORT", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=10, space_after=12, font_size=13, bold=True, font_color=RGBColor(70, 70, 70))
    p_title_top.paragraph_format.line_spacing = 1.15
    
    p_title = add_p(doc, "CONVERGENCE-AWARE ACTIVE RECONFIGURABLE INTELLIGENT SURFACE FOR OVER-THE-AIR FEDERATED LEARNING UNDER IMPERFECT CSI", 
                    align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=24, font_size=18, bold=True, font_color=RGBColor(15, 23, 42))
    p_title.paragraph_format.line_spacing = 1.25

    p_sub = add_p(doc, "Submitted in partial fulfillment of the requirements\nfor the award of the degree of", 
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24, space_after=12, font_size=12, italic=True)
    p_sub.paragraph_format.line_spacing = 1.2

    p_deg = add_p(doc, "BACHELOR OF ENGINEERING\nIN\nELECTRONICS AND COMMUNICATION ENGINEERING", 
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=28, font_size=14, bold=True, font_color=RGBColor(30, 41, 59))
    p_deg.paragraph_format.line_spacing = 1.2

    p_by = add_p(doc, "Submitted by\n\nRAJARAJAN PONRAJ\n(Reg. No.: 720721106000)", 
                 align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=36, font_size=13, bold=True)
    p_by.paragraph_format.line_spacing = 1.2

    p_inst = add_p(doc, "DEPARTMENT OF ELECTRONICS AND COMMUNICATION ENGINEERING\n"
                        "HINDUSTHAN COLLEGE OF ENGINEERING AND TECHNOLOGY\n"
                        "(An Autonomous Institution, Affiliated to Anna University, Chennai)\n"
                        "COIMBATORE - 641 032, TAMIL NADU, INDIA\n\n"
                        "APRIL 2026", 
                   align=WD_ALIGN_PARAGRAPH.CENTER, space_before=36, space_after=10, font_size=12, bold=True, font_color=RGBColor(15, 23, 42))
    p_inst.paragraph_format.line_spacing = 1.2

    doc.add_page_break()

    # -------------------------------------------------------------
    # 2. BONAFIDE CERTIFICATE
    # -------------------------------------------------------------
    add_p(doc, "HINDUSTHAN COLLEGE OF ENGINEERING AND TECHNOLOGY", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=4, font_size=14, bold=True, font_color=RGBColor(15, 23, 42))
    add_p(doc, "(An Autonomous Institution, Affiliated to Anna University, Chennai)\nCOIMBATORE - 641 032", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=18, font_size=11, italic=True)
    
    add_p(doc, "BONAFIDE CERTIFICATE", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=20, font_size=16, bold=True, font_color=RGBColor(15, 23, 42))

    cert_text = (
        "Certified that this project phase report titled \"CONVERGENCE-AWARE ACTIVE RECONFIGURABLE INTELLIGENT "
        "SURFACE FOR OVER-THE-AIR FEDERATED LEARNING UNDER IMPERFECT CSI\" is the bonafide work of "
        "RAJARAJAN PONRAJ who carried out the research project work under my supervision in partial "
        "fulfillment of the requirements for the award of the degree of Bachelor of Engineering in Electronics and "
        "Communication Engineering of Hindusthan College of Engineering and Technology, Coimbatore during the academic year 2025-2026."
    )
    add_p(doc, cert_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=12, space_after=40, font_size=12, line_spacing=1.5)

    tbl_sig = doc.add_table(rows=1, cols=2)
    tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_sig.rows[0].cells[0].width = Inches(3.0)
    tbl_sig.rows[0].cells[1].width = Inches(3.0)

    p_sig1 = tbl_sig.rows[0].cells[0].paragraphs[0]
    p_sig1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r1 = p_sig1.add_run("SUPERVISOR\n\n\n\n___________________________\nProject Supervisor\nDepartment of ECE\nHICET, Coimbatore - 641 032")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)
    r1.bold = True

    p_sig2 = tbl_sig.rows[0].cells[1].paragraphs[0]
    p_sig2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r2 = p_sig2.add_run("HEAD OF THE DEPARTMENT\n\n\n\n___________________________\nDr. P. F. Khaleelur Rahiman\nProfessor & Head, Dept. of ECE\nHICET, Coimbatore - 641 032")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(11)
    r2.bold = True

    add_p(doc, "\n\nSubmitted for the Project Phase Viva-Voce examination held on ____________________", space_before=30, space_after=40, font_size=11, italic=True)

    tbl_eval = doc.add_table(rows=1, cols=2)
    tbl_eval.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_eval.rows[0].cells[0].width = Inches(3.0)
    tbl_eval.rows[0].cells[1].width = Inches(3.0)
    tbl_eval.rows[0].cells[0].paragraphs[0].add_run("___________________________\nINTERNAL EXAMINER").font.name = "Times New Roman"
    tbl_eval.rows[0].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    tbl_eval.rows[0].cells[1].paragraphs[0].add_run("___________________________\nEXTERNAL EXAMINER").font.name = "Times New Roman"

    doc.add_page_break()

    # -------------------------------------------------------------
    # 3. DECLARATION
    # -------------------------------------------------------------
    add_p(doc, "DECLARATION", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=24, font_size=16, bold=True, font_color=RGBColor(15, 23, 42))

    decl_text = (
        "I hereby declare that the project phase report titled \"CONVERGENCE-AWARE ACTIVE RECONFIGURABLE INTELLIGENT "
        "SURFACE FOR OVER-THE-AIR FEDERATED LEARNING UNDER IMPERFECT CSI\" submitted to the Department of Electronics and "
        "Communication Engineering, Hindusthan College of Engineering and Technology, Coimbatore, is a record of original work "
        "done by me under the supervision and guidance of my Project Guide.\n\n"
        "I further declare that this report has not formed the basis for the award of any other Degree, Diploma, Associateship, "
        "Fellowship, or any other similar title in any other Institution or University."
    )
    add_p(doc, decl_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=12, space_after=48, font_size=12, line_spacing=1.5)

    p_dec_sig = add_p(doc, "Place: Coimbatore\nDate: 13-04-2026\n\n\n\n\nRAJARAJAN PONRAJ\n(Reg. No.: 720721106000)", 
                      align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=20, space_after=10, font_size=12, bold=True)
    p_dec_sig.paragraph_format.line_spacing = 1.25

    doc.add_page_break()

    # -------------------------------------------------------------
    # 4. ACKNOWLEDGEMENTS
    # -------------------------------------------------------------
    add_p(doc, "ACKNOWLEDGEMENTS", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=24, font_size=16, bold=True, font_color=RGBColor(15, 23, 42))

    ack_text = (
        "First and foremost, I express my profound gratitude to the Almighty for bestowing grace, wisdom, and perseverance "
        "throughout the execution of this project phase.\n\n"
        "I express my sincere thanks to the Management of Hindusthan College of Engineering and Technology for providing state-of-the-art "
        "laboratory facilities and high-performance computing environments that fostered an inspiring research atmosphere.\n\n"
        "I extend my heartfelt gratitude to our respected Principal for his visionary encouragement and administrative support throughout "
        "the academic program.\n\n"
        "I place on record my sincere thanks to Dr. P. F. Khaleelur Rahiman, Professor & Head, Department of Electronics and Communication "
        "Engineering, for his invaluable encouragement, constant guidance, and support throughout the project phase development.\n\n"
        "I am immensely indebted to my Project Supervisor for their patient mentoring, insightful technical suggestions, and critical scrutiny "
        "of the theoretical mathematical formulations, channel modeling principles, and federated learning simulations.\n\n"
        "I also thank all faculty members, technical laboratory assistants, and fellow classmates of the Department of Electronics and "
        "Communication Engineering for their cooperation and camaraderie. Finally, I dedicate this work to my beloved parents and family "
        "members whose selfless sacrifices and unwavering prayers have been my persistent source of strength."
    )
    add_p(doc, ack_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=10, space_after=10, font_size=12, line_spacing=1.5)

    doc.add_page_break()

    # -------------------------------------------------------------
    # 5. ABSTRACT
    # -------------------------------------------------------------
    add_p(doc, "ABSTRACT", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=20, font_size=16, bold=True, font_color=RGBColor(15, 23, 42))

    abs_p1 = (
        "Wireless Over-the-Air Federated Learning (AirFL) represents a transformative edge intelligence paradigm that coordinates distributed "
        "edge devices to train a centralized machine learning model while preserving data privacy. By leveraging the superposition property of "
        "wireless electromagnetic waves across a multiple-access channel, Over-the-Air Computation (AirComp) performs simultaneous analog aggregation "
        "of high-dimensional local gradient updates, reducing uplink transmission latency to a constant that is independent of client population. "
        "However, physical AirComp is severely compromised by deep channel fading, non-line-of-sight blockage, and channel estimation inaccuracies. "
        "Although passive Reconfigurable Intelligent Surfaces (RIS) have been introduced to create virtual line-of-sight propagation links, their "
        "cascaded channel structure suffers from the fundamental 'double-fading' (multiplicative path-loss) attenuation, necessitating impractically "
        "massive surface arrays to achieve meaningful link budget improvements."
    )
    add_p(doc, abs_p1, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=8, space_after=10, font_size=12, line_spacing=1.5)

    abs_p2 = (
        "Active RIS technology overcomes cascaded path loss by integrating reflection-type amplifiers into the surface elements, providing simultaneous "
        "phase manipulation and electromagnetic amplification. Nevertheless, active surface elements introduce internal thermal noise that is amplified "
        "and propagated toward the Access Point (AP). Crucially, existing active-RIS AirComp frameworks universally employ naive instantaneous Mean Squared "
        "Error (MSE) minimization based on estimated Channel State Information (CSI). In this project phase, we mathematically uncover and empirically "
        "demonstrate that naive MSE minimization fails catastrophically under imperfect CSI: because the objective omits CSI estimation error variance, "
        "the solver greedily drives active amplification to its hardware maximum, which severely inflates both thermal noise and channel estimation errors, "
        "corrupting aggregated gradient updates and destabilizing federated optimization."
    )
    add_p(doc, abs_p2, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=8, space_after=10, font_size=12, line_spacing=1.5)

    abs_p3 = (
        "To resolve this fundamental vulnerability, this project develops a novel Convergence-Aware Optimization Framework for active-RIS-assisted "
        "AirFL operating under imperfect CSI. By evaluating the conditional statistical expectation of aggregation MSE over complex Gaussian CSI estimation "
        "uncertainties, we derive an exact closed-form objective exposing an analytical regularization penalty proportional to active amplification power, "
        "transmit power, and CSI error variance. Using the L-smooth descent lemma, we establish a theoretical guarantee linking the derived objective directly "
        "to per-round global loss reduction bounds. Furthermore, by exploiting Zero-Forcing (ZF) precoding, transceiver scalers are derived in closed form, "
        "collapsing the joint search space to N+1 surface parameters solved via multi-start bound-constrained Quasi-Newton (L-BFGS-B) optimization with O(N^2) "
        "computational complexity, bypassing expensive Semidefinite Relaxation (SDR) methods."
    )
    add_p(doc, abs_p3, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=8, space_after=10, font_size=12, line_spacing=1.5)

    abs_p4 = (
        "Extensive Monte Carlo wireless simulations (across 500 random channel seeds) and distributed PyTorch federated learning experiments on MNIST and "
        "CIFAR-10 under Dirichlet non-IID data partitioning (alpha = 0.5) confirm the superiority of the proposed framework. The convergence-aware controller "
        "achieves up to an order-of-magnitude (10x) reduction in aggregation MSE under high CSI uncertainty, prevents amplification noise explosion, and accelerates "
        "FL convergence speed by 3x. On MNIST, the proposed framework reaches 97.8% (+-0.4%) test accuracy, matching the ideal noise-free benchmark, while naive "
        "MSE minimization plateaus at 92.4% (+-2.8%). On CIFAR-10, the proposed controller attains 62.1% accuracy compared to 48.2% for the baseline, validating "
        "its robustness for next-generation 6G edge intelligence networks."
    )
    add_p(doc, abs_p4, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=8, space_after=12, font_size=12, line_spacing=1.5)

    p_kw = add_p(doc, "Keywords: Active Reconfigurable Intelligent Surface (Active RIS), Over-the-Air Computation (AirComp), Federated Learning (FL), "
                      "Imperfect Channel State Information (CSI), Convergence-Aware Optimization, Edge Intelligence, 6G Wireless Networks.", 
                 space_before=12, space_after=10, font_size=11, bold=True)
    p_kw.paragraph_format.line_spacing = 1.25

    doc.add_page_break()

    # -------------------------------------------------------------
    # 6. TABLE OF CONTENTS
    # -------------------------------------------------------------
    add_p(doc, "TABLE OF CONTENTS", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=20, font_size=16, bold=True, font_color=RGBColor(15, 23, 42))

    toc_items = [
        ("BONAFIDE CERTIFICATE", "ii"),
        ("DECLARATION", "iii"),
        ("ACKNOWLEDGEMENTS", "iv"),
        ("ABSTRACT", "v"),
        ("LIST OF FIGURES", "viii"),
        ("LIST OF TABLES", "ix"),
        ("LIST OF ABBREVIATIONS AND ACRONYMS", "x"),
        ("LIST OF SYMBOLS", "xi"),
        ("CHAPTER 1: INTRODUCTION", "1"),
        ("  1.1 Background and Paradigm Shift in Edge Computing", "1"),
        ("  1.2 The Communication Bottleneck in Distributed Learning", "2"),
        ("  1.3 Over-the-Air Computation (AirComp) Principles", "3"),
        ("  1.4 Reconfigurable Intelligent Surfaces: Passive vs Active Paradigms", "4"),
        ("  1.5 Problem Statement: Vulnerabilities Under Imperfect CSI", "5"),
        ("  1.6 Project Objectives", "6"),
        ("  1.7 Scope and Organization of the Report", "6"),
        ("CHAPTER 2: LITERATURE SURVEY AND RELATED WORK", "8"),
        ("  2.1 Evolution of Edge AI and Federated Learning", "8"),
        ("  2.2 Analog Over-the-Air Aggregation in Wireless Networks", "9"),
        ("  2.3 RIS-Assisted Wireless Communications and Cascaded Path Loss", "10"),
        ("  2.4 Active RIS Architectures and Thermal Noise Amplification", "11"),
        ("  2.5 Channel Estimation and CSI Imperfection in Edge Networks", "12"),
        ("  2.6 Comparative Literature Summary Table", "13"),
        ("  2.7 Research Gaps and Problem Formulation Rationale", "15"),
        ("CHAPTER 3: SYSTEM MODEL AND MATHEMATICAL FORMULATION", "17"),
        ("  3.1 Network Topology and Physical Deployment Model", "17"),
        ("  3.2 Federated Learning Framework and Empirical Risk Minimization", "18"),
        ("  3.3 Channel Propagation, Path Loss, and Fading Models (Rician & Rayleigh)", "19"),
        ("  3.4 Active RIS Signal, Hardware, and Thermal Noise Model", "20"),
        ("  3.5 Analog Over-the-Air Signal Superposition and Receive Estimator", "21"),
        ("  3.6 True Instantaneous Aggregation Mean Squared Error (MSE)", "22"),
        ("  3.7 Mathematical Breakdown of Naive Instantaneous MSE-Minimization", "23"),
        ("  3.8 Formulation of Proposed Convergence-Aware Objective (J_t)", "24"),
        ("  3.9 Derivation of Analytical Regularization Penalty (Delta J_t)", "25"),
        ("  3.10 Theoretical Linkage: Descent Lemma and FL Convergence Guarantee", "26"),
        ("CHAPTER 4: OPTIMIZATION FRAMEWORK AND PROPOSED METHODOLOGY", "28"),
        ("  4.1 Joint Optimization Formulation", "28"),
        ("  4.2 Closed-Form Zero-Forcing Transceiver Precoding", "29"),
        ("  4.3 Optimal Receive Normalization Factor (c*)", "30"),
        ("  4.4 Dimensionality Reduction: Collapsing the Search Space to (N+1)", "31"),
        ("  4.5 Bound-Constrained Quasi-Newton Surface Optimization (L-BFGS-B / SLSQP)", "32"),
        ("  4.6 Multi-Start Random Restart Safeguard Mechanism", "33"),
        ("  4.7 Complete Algorithmic Workflow and Structured Pseudo-Code", "34"),
        ("  4.8 Computational Complexity Analysis: Big-O Comparison Against SDR", "35"),
        ("  4.9 Real-Time Feasibility and Coherence Time Budgeting", "36"),
        ("CHAPTER 5: SIMULATION ARCHITECTURE AND EXPERIMENTAL SETUP", "37"),
        ("  5.1 End-to-End Simulation Pipeline and Modular Software Architecture", "37"),
        ("  5.2 Channel Generation and Imperfect CSI Synthesizer Modules", "38"),
        ("  5.3 Active and Passive RIS Simulation Modules", "39"),
        ("  5.4 AirComp Analog Signal Superposition and Noise Injection Engine", "40"),
        ("  5.5 PyTorch Distributed Federated Learning Pipeline", "41"),
        ("  5.6 Dataset Preprocessing and Non-IID Dirichlet Partitioning", "42"),
        ("  5.7 Comprehensive System and Simulation Parameter Calibration", "43"),
        ("CHAPTER 6: RESULTS, PERFORMANCE EVALUATION AND DISCUSSION", "45"),
        ("  6.1 Baseline Schemes for Performance Comparison", "45"),
        ("  6.2 Physical Layer Wireless Performance Evaluations", "46"),
        ("    6.2.1 Exp 1: Expected MSE vs. Amplification Factor Trade-Off", "46"),
        ("    6.2.2 Exp 2: Expected MSE vs. Transmit Power Budget / SNR", "47"),
        ("    6.2.3 Exp 3: Optimal Amplification Factor Adaptation vs. SNR", "48"),
        ("    6.2.4 Exp 4: Robustness Under Imperfect CSI Variance", "49"),
        ("    6.2.5 Exp 5: RIS Element Count Array Scalability (N)", "50"),
        ("    6.2.6 Exp 6: Edge Client Population Scalability (K)", "51"),
        ("  6.3 End-to-End Federated Learning Performance Evaluations", "52"),
        ("    6.3.1 Exp 7 & 8: FL Convergence Under IID vs. Non-IID Partitioning", "52"),
        ("    6.3.2 Exp 9: Monte Carlo Statistical Confidence Over 500 Seeds", "53"),
        ("    6.3.3 Exp 10: Benchmark on Complex Vision Dataset (CIFAR-10)", "54"),
        ("  6.4 Quantitative Performance Summary Table Across All Baselines", "55"),
        ("  6.5 Critical Engineering Insights and Trade-Off Analysis", "56"),
        ("CHAPTER 7: CONCLUSION AND FUTURE SCOPE", "58"),
        ("  7.1 Summary of Project Phase Work", "58"),
        ("  7.2 Key Contributions and Technical Milestones", "59"),
        ("  7.3 Practical Limitations and Hardware Constraints", "60"),
        ("  7.4 Directions for Project Phase II and Future Research", "61"),
        ("REFERENCES", "63"),
        ("APPENDIX A: SYSTEM PARAMETER REGISTER & NOTATION AUDIT", "67"),
        ("APPENDIX B: CORE SIMULATION ENGINE IMPLEMENTATION EXCERPTS", "70"),
    ]

    tbl_toc = doc.add_table(rows=len(toc_items), cols=2)
    tbl_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_toc, color="E2E8F0", sz="2")
    for idx, (title, page) in enumerate(toc_items):
        cell_t = tbl_toc.rows[idx].cells[0]
        cell_p = tbl_toc.rows[idx].cells[1]
        cell_t.width = Inches(5.4)
        cell_p.width = Inches(0.8)
        set_cell_margins(cell_t, top=40, bottom=40, left=60, right=60)
        set_cell_margins(cell_p, top=40, bottom=40, left=60, right=60)
        
        p_t = cell_t.paragraphs[0]
        p_t.paragraph_format.space_before = Pt(0)
        p_t.paragraph_format.space_after = Pt(0)
        p_t.paragraph_format.line_spacing = 1.15
        r_t = p_t.add_run(title)
        r_t.font.name = "Times New Roman"
        r_t.font.size = Pt(10.5)
        if title.startswith("CHAPTER") or title in ["BONAFIDE CERTIFICATE", "DECLARATION", "ACKNOWLEDGEMENTS", "ABSTRACT", "REFERENCES"]:
            r_t.bold = True
            r_t.font.color.rgb = RGBColor(15, 23, 42)
            
        p_p = cell_p.paragraphs[0]
        p_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_p.paragraph_format.space_before = Pt(0)
        p_p.paragraph_format.space_after = Pt(0)
        r_p = p_p.add_run(page)
        r_p.font.name = "Times New Roman"
        r_p.font.size = Pt(10.5)
        if title.startswith("CHAPTER"):
            r_p.bold = True

    doc.add_page_break()

    # -------------------------------------------------------------
    # 7. LIST OF FIGURES & TABLES
    # -------------------------------------------------------------
    add_p(doc, "LIST OF FIGURES", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=18, font_size=16, bold=True, font_color=RGBColor(15, 23, 42))

    fig_items = [
        ("Figure 3.1", "System architecture of active-RIS-assisted Over-the-Air Federated Learning under imperfect CSI", "18"),
        ("Figure 6.1", "Expected MSE vs. maximum amplification factor a_max under moderate CSI error", "46"),
        ("Figure 6.2", "Expected MSE vs. client transmit power budget P_max (SNR sweep)", "47"),
        ("Figure 6.3", "Optimal amplification factor adaptation a* vs. transmit power budget P_max", "48"),
        ("Figure 6.4", "Expected MSE vs. CSI estimation error variance sigma_e^2", "49"),
        ("Figure 6.5", "Expected MSE vs. active RIS reflecting element count N", "50"),
        ("Figure 6.6", "Expected MSE vs. edge client population K (bottleneck scaling)", "51"),
        ("Figure 6.7", "Top-1 test accuracy vs. communication rounds on MNIST under non-IID Dirichlet partitioning", "53"),
        ("Figure 6.8", "Top-1 test accuracy vs. communication rounds on CIFAR-10 under non-IID Dirichlet partitioning", "54"),
    ]

    tbl_lof = doc.add_table(rows=len(fig_items), cols=3)
    tbl_lof.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_lof, color="E2E8F0", sz="2")
    for idx, (f_num, f_title, f_page) in enumerate(fig_items):
        c0, c1, c2 = tbl_lof.rows[idx].cells[0], tbl_lof.rows[idx].cells[1], tbl_lof.rows[idx].cells[2]
        c0.width, c1.width, c2.width = Inches(1.2), Inches(4.2), Inches(0.8)
        set_cell_margins(c0, 40, 40, 60, 60)
        set_cell_margins(c1, 40, 40, 60, 60)
        set_cell_margins(c2, 40, 40, 60, 60)
        
        r0 = c0.paragraphs[0].add_run(f_num)
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(10.5)
        r0.bold = True
        
        r1 = c1.paragraphs[0].add_run(f_title)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(10.5)
        
        c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r2 = c2.paragraphs[0].add_run(f_page)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(10.5)

    add_p(doc, "LIST OF TABLES", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=24, space_after=18, font_size=16, bold=True, font_color=RGBColor(15, 23, 42))

    tab_items = [
        ("Table 2.1", "Comprehensive comparison of state-of-the-art RIS-assisted AirComp-FL frameworks", "14"),
        ("Table 4.1", "Asymptotic computational complexity comparison across optimization paradigms", "35"),
        ("Table 5.1", "System parameters, channel assumptions, and federated learning hyperparameter register", "43"),
        ("Table 6.1", "Consolidated performance evaluation matrix across all benchmarks and metrics", "56"),
        ("Table A.1", "Mathematical symbol inventory, physical dimensions, and operational definitions", "67"),
    ]

    tbl_lot = doc.add_table(rows=len(tab_items), cols=3)
    tbl_lot.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_lot, color="E2E8F0", sz="2")
    for idx, (t_num, t_title, t_page) in enumerate(tab_items):
        c0, c1, c2 = tbl_lot.rows[idx].cells[0], tbl_lot.rows[idx].cells[1], tbl_lot.rows[idx].cells[2]
        c0.width, c1.width, c2.width = Inches(1.2), Inches(4.2), Inches(0.8)
        set_cell_margins(c0, 40, 40, 60, 60)
        set_cell_margins(c1, 40, 40, 60, 60)
        set_cell_margins(c2, 40, 40, 60, 60)
        
        r0 = c0.paragraphs[0].add_run(t_num)
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(10.5)
        r0.bold = True
        
        r1 = c1.paragraphs[0].add_run(t_title)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(10.5)
        
        c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r2 = c2.paragraphs[0].add_run(t_page)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(10.5)

    doc.add_page_break()

    # -------------------------------------------------------------
    # 8. LIST OF ABBREVIATIONS AND SYMBOLS
    # -------------------------------------------------------------
    add_p(doc, "LIST OF ABBREVIATIONS AND ACRONYMS", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=16, font_size=16, bold=True, font_color=RGBColor(15, 23, 42))

    abbrs = [
        ("3GPP", "3rd Generation Partnership Project"),
        ("6G", "Sixth-Generation Cellular Wireless Networks"),
        ("AirComp", "Over-the-Air Computation"),
        ("AirFL", "Over-the-Air Federated Learning"),
        ("AO", "Alternating Optimization"),
        ("AP", "Access Point (Parameter Server)"),
        ("AWGN", "Additive White Gaussian Noise"),
        ("CSI", "Channel State Information"),
        ("FL", "Federated Learning"),
        ("IID", "Independent and Identically Distributed"),
        ("L-BFGS-B", "Limited-memory Broyden-Fletcher-Goldfarb-Shanno Bound-constrained"),
        ("LoS", "Line-of-Sight"),
        ("MAC", "Multiple Access Channel"),
        ("MSE", "Mean Squared Error"),
        ("NLoS", "Non-Line-of-Sight"),
        ("Non-IID", "Non-Independent and Identically Distributed"),
        ("OMA", "Orthogonal Multiple Access"),
        ("RF", "Radio Frequency"),
        ("RIS", "Reconfigurable Intelligent Surface"),
        ("SDR", "Semidefinite Relaxation"),
        ("SGD", "Stochastic Gradient Descent"),
        ("SLSQP", "Sequential Least Squares Programming"),
        ("SNR", "Signal-to-Noise Ratio"),
        ("ZF", "Zero-Forcing"),
    ]

    tbl_abbr = doc.add_table(rows=len(abbrs), cols=2)
    tbl_abbr.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_abbr, color="E2E8F0", sz="2")
    for idx, (ab, exp) in enumerate(abbrs):
        c0, c1 = tbl_abbr.rows[idx].cells[0], tbl_abbr.rows[idx].cells[1]
        c0.width, c1.width = Inches(1.8), Inches(4.4)
        set_cell_margins(c0, 30, 30, 60, 60)
        set_cell_margins(c1, 30, 30, 60, 60)
        r0 = c0.paragraphs[0].add_run(ab)
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(10)
        r0.bold = True
        r1 = c1.paragraphs[0].add_run(exp)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(10)

    doc.add_page_break()

    add_p(doc, "LIST OF SYMBOLS", align=WD_ALIGN_PARAGRAPH.CENTER, space_before=18, space_after=16, font_size=16, bold=True, font_color=RGBColor(15, 23, 42))

    symbols = [
        ("K", "Total number of participating edge clients (user population)"),
        ("N", "Total number of reflecting elements on the active RIS array"),
        ("d", "Dimension of the machine learning model parameter/gradient vector"),
        ("w_t", "Global model weight parameter vector at communication round t"),
        ("g_{k,t}", "Local gradient update vector calculated by client k at round t"),
        ("p_k, alpha_k", "Statistical aggregation weight associated with client k"),
        ("h_k", "Uplink channel vector from edge client k to the active RIS (N x 1)"),
        ("G^H", "Uplink channel response vector from active RIS array to AP (1 x N)"),
        ("h_{d,k}", "Direct wireless channel scalar from client k to the AP receiver"),
        ("h_k^{eff}", "Effective composite channel response scalar for client k"),
        ("e_k", "Channel estimation error vector corrupting channel h_k"),
        ("sigma_e^2", "Statistical variance of channel state estimation error"),
        ("Theta", "Diagonal reflection and phase-shift matrix of the active RIS array"),
        ("a", "Active RIS uniform reflection amplification factor (voltage gain)"),
        ("a_max", "Hardware ceiling for active amplification factor"),
        ("theta_n", "Continuous phase shift imparted by RIS reflecting element n"),
        ("n_R", "Inherent thermal noise vector generated by active RIS circuitry"),
        ("sigma_R^2", "Active RIS internal thermal noise variance (Watts / dBm)"),
        ("n_0", "Background Additive White Gaussian Noise (AWGN) at the AP receiver"),
        ("sigma_0^2", "AP receiver background AWGN power variance (Watts / dBm)"),
        ("b_k", "Complex transmit precoding/scaling scalar applied by client k"),
        ("c", "Receive normalization scalar applied at the central AP receiver"),
        ("P_max", "Maximum individual transmit power budget constraint per edge client"),
        ("P_{RIS,max}", "Total power budget allocated for active RIS surface amplification"),
        ("J_t", "Proposed Convergence-Aware expected aggregation MSE objective"),
        ("Delta J_t", "Derived analytical regularization penalty term enforcing gain back-off"),
        ("L", "Lipschitz smoothness constant of the global federated loss function"),
        ("eta", "Global learning rate step size applied during server weight updates"),
    ]

    tbl_sym = doc.add_table(rows=len(symbols), cols=2)
    tbl_sym.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_sym, color="E2E8F0", sz="2")
    for idx, (sym, desc) in enumerate(symbols):
        c0, c1 = tbl_sym.rows[idx].cells[0], tbl_sym.rows[idx].cells[1]
        c0.width, c1.width = Inches(1.8), Inches(4.4)
        set_cell_margins(c0, 30, 30, 60, 60)
        set_cell_margins(c1, 30, 30, 60, 60)
        r0 = c0.paragraphs[0].add_run(sym)
        r0.font.name = "Times New Roman"
        r0.font.size = Pt(10)
        r0.bold = True
        r0.italic = True
        r1 = c1.paragraphs[0].add_run(desc)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(10)

    doc.add_page_break()
