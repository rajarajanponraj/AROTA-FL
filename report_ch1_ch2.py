import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx_helpers import (
    add_p, add_heading_1, add_heading_2, add_heading_3, 
    set_cell_background, set_cell_margins, set_table_borders, add_callout
)

def build_chapter_1(doc):
    # -------------------------------------------------------------
    # CHAPTER 1: INTRODUCTION
    # -------------------------------------------------------------
    add_heading_1(doc, "CHAPTER 1\nINTRODUCTION", space_before=12, space_after=14)

    add_heading_2(doc, "1.1 Background and Paradigm Shift in Edge Computing")
    p1 = (
        "The proliferation of ubiquitous Internet of Things (IoT) devices, autonomous vehicular systems, and smart edge sensors has driven an "
        "unprecedented surge in the volume of localized data generated at the network periphery. Traditional centralized machine learning architectures "
        "mandate the continuous backhauling of raw training data from geographically distributed clients to a centralized cloud parameter server. "
        "However, this centralized paradigm encounters severe operational bottlenecks, including prohibitive uplink bandwidth consumption, excessive "
        "transmission latency, and critical privacy vulnerabilities associated with transmitting proprietary personal data across untrusted public networks. "
        "To mitigate these challenges, Federated Learning (FL) has emerged as a groundbreaking decentralized machine learning framework that enables edge "
        "devices to collaboratively train a shared global model while strictly retaining raw datasets locally on device storage."
    )
    add_p(doc, p1)

    p2 = (
        "In a canonical federated learning iteration, edge clients download the latest global model parameters, compute local stochastic gradient estimates "
        "over their private local datasets, and periodically transmit only their model updates (or gradient vectors) to an Access Point (AP) or base station. "
        "The AP aggregates these distributed client updates—typically computing a weighted average—and broadcasts the updated global model back to edge devices "
        "for subsequent optimization rounds. By eliminating the necessity of transmitting raw user samples, federated learning guarantees data confidentiality, "
        "complies with strict regulatory data protection standards, and significantly reduces raw data uplink payloads."
    )
    add_p(doc, p2)

    add_heading_2(doc, "1.2 The Communication Bottleneck in Distributed Learning")
    p3 = (
        "Despite its privacy and data governance advantages, the scalability of conventional wireless federated learning is fundamentally constrained "
        "by uplink communication overhead. State-of-the-art Deep Neural Networks (DNNs) frequently comprise tens of millions of trainable parameters. "
        "In traditional digital wireless implementations based on Orthogonal Multiple Access (OMA) mechanisms—such as Time Division Multiple Access (TDMA), "
        "Frequency Division Multiple Access (FDMA), or Orthogonal Frequency Division Multiple Access (OFDMA)—participating edge devices must be allocated "
        "disjoint orthogonal time slots or frequency sub-bands to upload their high-dimensional parameter updates. Consequently, the required total uplink "
        "bandwidth and communication latency scale linearly with the participating client population K."
    )
    add_p(doc, p3)

    p4 = (
        "When deployed across dense 6G edge environments where hundreds of edge devices simultaneously participate in training, orthogonal radio resource "
        "slicing leads to severe spectrum congestion, excessive channel latency, and communication straggler effects, where the slowest wireless links dictate "
        "the per-round synchronization delay. As a consequence, communication latency overwhelmingly dominates computational execution time, frequently "
        "accounting for over 80% of total training turnaround time in real-world distributed wireless deployments."
    )
    add_p(doc, p4)

    add_heading_2(doc, "1.3 Over-the-Air Computation (AirComp) Principles")
    p5 = (
        "To break the linear latency scaling law inherent in orthogonal communications, Over-the-Air Computation (AirComp) has been proposed as a revolutionary "
        "analog physical-layer aggregation mechanism. Rather than decoupling digital transmission and parameter aggregation into separate processing steps, "
        "AirComp seamlessly integrates wireless transmission with mathematical computation by exploiting the natural superposition property of electromagnetic "
        "waves across an uncoordinated multiple-access radio channel (MAC)."
    )
    add_p(doc, p5)

    p6 = (
        "Under AirComp-assisted Federated Learning (AirFL), all participating edge devices synchronize their local transmissions and simultaneously upload "
        "their analog gradient symbols across the exact same frequency band. The multiple-access channel acts as an analog computing engine, such that the "
        "signal received at the AP is the physical sum of all individual waveforms. In a single channel use, the AP directly receives the aggregated global update, "
        "collapsing the required aggregation latency to a constant independent of the client count K. This paradigm shift provides an immense spectrum efficiency "
        "gain and unlocks massive ultra-dense edge intelligence."
    )
    add_p(doc, p6)

    p7 = (
        "However, the operational success of analog AirComp hinges upon precise phase synchronization and strict amplitude alignment among all participating "
        "devices. Because analog signals undergo destructive interference if misaligned, deep channel fading, Non-Line-of-Sight (NLoS) blockage, or severe "
        "propagation attenuation severely corrupts the received superposition. In particular, AirComp is inherently vulnerable to the 'channel bottleneck effect': "
        "to ensure unbiased model aggregation, edge devices must invert their channel fading via pre-equalization scaling. When edge clients experience deep "
        "fading or severe shadowing, the weakest client dictates the receive normalization factor, which substantially amplifies receiver background noise and "
        "injects severe aggregation errors into the global machine learning update."
    )
    add_p(doc, p7)

    add_heading_2(doc, "1.4 Reconfigurable Intelligent Surfaces: Passive vs Active Paradigms")
    p8 = (
        "To establish favorable propagation conditions and eliminate channel bottlenecks, Reconfigurable Intelligent Surfaces (RIS) have recently emerged as "
        "a cornerstone physical-layer enhancement for 6G wireless systems. An RIS consists of an array of sub-wavelength electromagnetic metasurface elements "
        "capable of dynamically altering the phase, amplitude, or polarization of incident radio frequency (RF) waves under software control, thereby shaping the "
        "wireless environment into a smart, programmable entity."
    )
    add_p(doc, p8)

    p9 = (
        "Traditional RIS architectures are strictly passive, comprising PIN diodes, varactors, or liquid crystals that adjust reflection phases without active RF "
        "amplification circuits. While passive RIS boasts near-zero power consumption, it suffers from a devastating fundamental propagation penalty known as "
        "'double-fading' (cascaded multiplicative path loss). Specifically, the signal traveling from an edge user to the RIS and subsequently from the RIS to "
        "the AP undergoes path loss proportional to the product of both distance components (PL ~ d_1^alpha1 * d_2^alpha2). Consequently, passive surfaces require "
        "thousands of reflecting elements to overcome cascaded propagation losses, resulting in massive physical apertures and prohibitive beamforming complexity."
    )
    add_p(doc, p9)

    p10 = (
        "To conquer the double-fading barrier, Active RIS technology has emerged as a disruptive breakthrough. Active RIS arrays integrate reflection-type "
        "RF amplifiers (e.g., tunnel diodes or low-noise operational amplifiers) directly into each reflecting element. By simultaneously adjusting the phase shift "
        "and amplifying the reflected wave's amplitude, an active RIS overcomes cascaded attenuation using compact physical aperture dimensions. However, this active "
        "amplification fundamentally changes the physical-layer noise dynamics: active amplifiers inevitably generate internal hardware thermal noise at the surface, "
        "which is amplified and forwarded to the central AP alongside desired gradient signals."
    )
    add_p(doc, p10)

    add_heading_2(doc, "1.5 Problem Statement: Vulnerabilities Under Imperfect CSI")
    p11 = (
        "In realistic 6G millimeter-wave and sub-terahertz edge deployments, acquiring perfect Channel State Information (CSI) for high-dimensional cascaded RIS "
        "links is practically impossible due to pilot contamination, channel estimation errors, and rapid channel aging. In existing literature, beamforming "
        "and active RIS configurations are designed by minimizing the estimated instantaneous aggregation Mean Squared Error (MSE), treating imperfect channel "
        "estimates as absolute ground truth."
    )
    add_p(doc, p11)

    p12 = (
        "This project phase identifies a fundamental, previously unaddressed failure mode in existing active-RIS AirFL designs: under imperfect CSI, naive "
        "instantaneous MSE minimization breaks down catastrophically. Because the objective function is computed purely on erroneous channel estimates, it fails "
        "to penalize the stochastic variance of channel estimation errors. Consequently, the optimization algorithm aggressively drives active amplification gains "
        "to their maximum hardware ceiling, believing that higher amplification monotonically suppresses receiver noise. In reality, this excessive amplification "
        "heavily multiplies both the surface thermal noise and the underlying channel estimation errors, injecting severe variance into the global gradient update, "
        "destabilizing federated learning convergence, and degrading final classification accuracy."
    )
    add_p(doc, p12)

    add_callout(doc, "Core Research Hypothesis", 
                "Under imperfect CSI and active amplification noise, the physical-layer configuration that minimizes naive estimated instantaneous MSE is "
                "fundamentally misaligned with the objective of maximizing end-to-end federated learning convergence. Formulating a convergence-aware "
                "objective that explicitly accounts for CSI error statistics will naturally regularize amplification gain, preventing noise divergence and "
                "guaranteeing monotonic learning progress.")

    add_heading_2(doc, "1.6 Project Objectives")
    p_obj = (
        "To resolve these critical vulnerabilities, this project phase establishes the following rigorous engineering objectives:\n\n"
        "1. Mathematical Formulation: Formulate an exact conditional statistical expectation model of aggregation MSE over complex Gaussian CSI estimation "
        "uncertainties, uncovering the analytical regularization penalty that governs active amplification gain.\n\n"
        "2. Algorithmic Design: Develop a computationally efficient joint optimization framework that determines closed-form Zero-Forcing (ZF) transceiver "
        "scalers and optimizes active surface phases and gains via Quasi-Newton methods with O(N^2) complexity, bypassing expensive Semidefinite Relaxation (SDR).\n\n"
        "3. Cross-Layer Simulation Architecture: Construct a modular, end-to-end Python/PyTorch simulation suite integrating realistic 3GPP distance-dependent "
        "path loss, Rician/Rayleigh fading, imperfect CSI generation, Active RIS hardware constraints, and distributed federated optimization.\n\n"
        "4. Empirical Validation and Benchmarking: Conduct extensive Monte Carlo physical-layer sweeps (500 random channel seeds) and federated learning training "
        "benchmarks on MNIST and CIFAR-10 under severe Dirichlet non-IID data partitioning (alpha = 0.5), validating convergence speed, robustness to CSI error, "
        "and final accuracy improvements against standard baselines."
    )
    add_p(doc, p_obj)

    add_heading_2(doc, "1.7 Scope and Organization of the Report")
    p_org = (
        "This project phase report documents the complete theoretical modeling, algorithmic design, implementation pipeline, and experimental performance "
        "evaluation of the proposed Convergence-Aware Active RIS AirFL framework. The remainder of this report is organized as follows:\n\n"
        "• Chapter 2 provides an exhaustive literature survey reviewing edge AI, AirComp aggregation, passive and active RIS models, and imperfect CSI estimation, "
        "highlighting the identified research gaps through a detailed comparative matrix.\n\n"
        "• Chapter 3 presents the formal system model, channel fading formulations, active RIS hardware noise dynamics, and derives the proposed Convergence-Aware "
        "expected MSE objective alongside the theoretical descent lemma proof.\n\n"
        "• Chapter 4 details the joint optimization methodology, zero-forcing transceiver derivations, bound-constrained Quasi-Newton algorithm, and Big-O computational "
        "complexity analysis.\n\n"
        "• Chapter 5 describes the software implementation architecture, modular codebase structure, PyTorch federated learning engine, and simulation calibration.\n\n"
        "• Chapter 6 presents comprehensive simulation results spanning wireless physical-layer sweeps, Monte Carlo confidence intervals, and distributed FL convergence "
        "trajectories on vision datasets, accompanied by critical engineering discussions.\n\n"
        "• Chapter 7 concludes the report with a synthesis of major achievements, deployment considerations, and directions for Project Phase II."
    )
    add_p(doc, p_org)

    doc.add_page_break()

def build_chapter_2(doc):
    # -------------------------------------------------------------
    # CHAPTER 2: LITERATURE SURVEY AND RELATED WORK
    # -------------------------------------------------------------
    add_heading_1(doc, "CHAPTER 2\nLITERATURE SURVEY AND RELATED WORK", space_before=12, space_after=14)

    add_heading_2(doc, "2.1 Evolution of Edge AI and Federated Learning")
    p1 = (
        "The convergence of edge computing and artificial intelligence has established Edge AI as a transformative paradigm for next-generation 6G mobile "
        "networks. In their seminal survey, Lim et al. (2020) comprehensively examined the architectural principles of mobile edge federated learning, categorizing "
        "resource allocation strategies, incentive mechanisms, and privacy preservation protocols. They demonstrated that while federated optimization protects "
        "user data locality, communication latency represents the primary bottleneck hindering practical edge deployment, especially across bandwidth-constrained "
        "cellular networks. Wang et al. (2019) investigated adaptive federated learning in resource-constrained wireless systems, proposing dynamic local epoch "
        "control to balance local computation energy against wireless transmission overhead. However, both foundational works operated under conventional "
        "orthogonal digital transmission models, inheriting the linear delay scaling penalty as edge client counts increase."
    )
    add_p(doc, p1)

    add_heading_2(doc, "2.2 Analog Over-the-Air Aggregation in Wireless Networks")
    p2 = (
        "To overcome the communication bottlenecks of orthogonal multiple access, researchers turned to analog Over-the-Air Computation (AirComp). Zhu et al. (2021) "
        "provided a comprehensive architectural review of AirComp for massive IoT data aggregation, demonstrating that analog signal superposition effectively "
        "computes nomographic functions directly through physical electromagnetic wave interference. Cao et al. (2020) formulated optimized transmit power control "
        "schemes for AirComp across fading channels, proving that channel-inversion pre-equalization minimizes aggregation mean squared error. In subsequent work, "
        "Cao et al. (2022) extended power control specifically to Over-the-Air Federated Averaging (AirFL), rigorously proving that physical-layer aggregation error "
        "variance acts as an additive gradient perturbation that slows down the convergence of stochastic gradient descent. Zhu et al. (2024) and Liu et al. (2020) "
        "further analyzed AirFL scaling laws, demonstrating that while AirComp achieves order-of-magnitude latency reductions, its performance deteriorates rapidly "
        "when wireless channels suffer from deep fading or severe non-line-of-sight propagation blockage."
    )
    add_p(doc, p2)

    add_heading_2(doc, "2.3 RIS-Assisted Wireless Communications and Cascaded Path Loss")
    p3 = (
        "To mitigate propagation blockage and enhance channel quality, Reconfigurable Intelligent Surfaces (RIS) have been widely integrated into wireless "
        "networks. Yang et al. (2020) introduced passive RIS arrays into edge machine learning, demonstrating that intelligent phase shifting can align disparate "
        "user channels to boost federated learning performance. Ni et al. (2021) developed a multi-RIS-aided federated learning architecture, jointly optimizing "
        "client device selection and surface phase shifts. Zheng et al. (2022) investigated the trade-off between aggregation accuracy and data integrity in "
        "RIS-assisted AirFL, deriving alternating optimization algorithms to minimize aggregation MSE."
    )
    add_p(doc, p3)

    p4 = (
        "However, as highlighted by Zhang et al. (2023), passive RIS architectures are fundamentally impaired by 'double-fading' (multiplicative cascaded path loss). "
        "Because the incident RF signal undergoes path loss across both the transmitter-to-RIS link and the RIS-to-receiver link without amplification, the composite "
        "link gain scales inversely with the product of distances (PL ~ d_1^alpha1 * d_2^alpha2). Consequently, passive surfaces provide negligible signal enhancement "
        "unless equipped with hundreds or thousands of elements, introducing massive physical hardware footprints and severe channel estimation overhead."
    )
    add_p(doc, p4)

    add_heading_2(doc, "2.4 Active RIS Architectures and Thermal Noise Amplification")
    p5 = (
        "To overcome the severe cascaded path loss of passive surfaces, Active RIS technology was introduced by Long et al. (2021). By integrating negative "
        "resistance reflection amplifiers into surface elements, active RIS arrays independently control reflection phases while actively amplifying signal "
        "voltages. Zhang et al. (2023) conducted a comprehensive comparative analysis between active and passive RIS architectures in 6G communications, proving "
        "that active RIS can achieve up to a 20 dB link budget enhancement over passive RIS under identical physical array sizes. Yu et al. (2024) demonstrated the "
        "efficacy of active RIS in integrated sensing and communication (ISAC) systems, highlighting its ability to overcome severe path losses in millimeter-wave "
        "bands."
    )
    add_p(doc, p5)

    p6 = (
        "Nevertheless, the operational advantages of active RIS are accompanied by a critical physical challenge: active amplifier circuits inherently inject "
        "thermal noise at each surface element. As established by Long et al. (2021), this surface thermal noise is amplified by the active reflection gain and "
        "propagated to the central receiver. When multiple clients transmit simultaneously in AirComp, the amplified surface noise superimposes with background "
        "receiver AWGN, creating a delicate trade-off between signal amplification and thermal noise inflation."
    )
    add_p(doc, p6)

    add_heading_2(doc, "2.5 Channel Estimation and CSI Imperfection in Edge Networks")
    p7 = (
        "Acquiring high-quality Channel State Information (CSI) is vital for physical-layer pre-equalization in AirComp. However, in high-frequency cascaded RIS "
        "systems, obtaining accurate CSI is exceptionally difficult due to the passive or quasi-passive nature of reflecting arrays. Zhou et al. (2024) reviewed "
        "machine learning and heuristic optimization techniques for RIS-assisted networks, noting that channel estimation errors remain the primary bottleneck in "
        "practical deployments. Yao et al. (2023) analytically modeled imperfect CSI in AirFL, proving that channel estimation errors induce signal misalignment and "
        "introduce an irreducible error floor in model aggregation. Yang et al. (2025) and Pradhan et al. (2024) explored robust transmission designs for active "
        "RIS under imperfect CSI, demonstrating that conventional designs that treat estimated channels as perfect truth suffer severe performance collapse."
    )
    add_p(doc, p7)

    add_heading_2(doc, "2.6 Comparative Literature Summary Table")
    p8 = "Table 2.1 provides an exhaustive comparative taxonomy of landmark contributions in the domain of RIS-assisted AirComp and federated learning."
    add_p(doc, p8)

    # 12-paper comparison table
    comp_data = [
        ("Reference & Year", "System Paradigm", "RIS Architecture", "CSI Assumption", "Optimization Objective", "Key Limitation"),
        ("Lim et al. (2020)", "Standard FL", "None (No RIS)", "Perfect CSI", "Communication-Efficient FL", "OMA delay scales linearly O(K)"),
        ("Wang et al. (2019)", "Edge FL", "None (No RIS)", "Perfect CSI", "Resource-Adaptive SGD", "Orthogonal bandwidth bottleneck"),
        ("Zhu et al. (2021)", "AirComp IoT", "None (No RIS)", "Perfect CSI", "MSE Minimization", "Severe channel bottleneck under fading"),
        ("Cao et al. (2022)", "AirFL Edge", "None (No RIS)", "Perfect CSI", "Power Control for AirFL", "Vulnerable to severe NLoS blockage"),
        ("Yang et al. (2020)", "AirFL + RIS", "Passive RIS", "Perfect CSI", "MSE Minimization", "Double-fading limits coverage"),
        ("Ni et al. (2021)", "Multi-RIS FL", "Passive RIS", "Perfect CSI", "Alternating Optimization", "Cascaded path loss requires massive N"),
        ("Zheng et al. (2022)", "AirFL + RIS", "Passive RIS", "Imperfect CSI", "Accuracy-Integrity Trade-off", "Passive gain insufficient for edge users"),
        ("Long et al. (2021)", "Cellular Comm", "Active RIS", "Perfect CSI", "SNR Maximization", "Omits federated learning dynamics"),
        ("Zhang et al. (2023)", "6G Cellular", "Active vs Passive", "Perfect CSI", "Capacity Scaling Laws", "Assumes ideal error-free channel knowledge"),
        ("Yao et al. (2023)", "AirFL", "None (No RIS)", "Imperfect CSI", "Robust Transmit Power", "Cannot overcome severe propagation blockage"),
        ("Yang et al. (2025)", "Active RIS", "Active RIS", "Imperfect CSI", "Robust Beamforming", "Point-to-point only; omits AirComp & FL"),
        ("Proposed (2026)", "AirFL + Active RIS", "Active RIS", "Imperfect CSI", "Convergence-Aware Expected MSE", "None (Robust O(N^2) learning-aligned design)"),
    ]

    tbl_lit = doc.add_table(rows=len(comp_data), cols=6)
    tbl_lit.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_lit, color="CBD5E1", sz="3")
    
    col_widths = [Inches(1.1), Inches(1.0), Inches(0.9), Inches(0.9), Inches(1.3), Inches(1.2)]
    for r_idx, row in enumerate(tbl_lit.rows):
        is_hdr = (r_idx == 0)
        is_prop = (r_idx == len(comp_data) - 1)
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths[c_idx]
            set_cell_margins(cell, top=60, bottom=60, left=70, right=70)
            if is_hdr:
                set_cell_background(cell, "1E293B")
            elif is_prop:
                set_cell_background(cell, "EFF6FF")
            elif r_idx % 2 == 1:
                set_cell_background(cell, "F8FAFC")
                
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(comp_data[r_idx][c_idx])
            run.font.name = "Times New Roman"
            run.font.size = Pt(8.5)
            if is_hdr:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            elif is_prop:
                run.bold = True
                run.font.color.rgb = RGBColor(37, 99, 235)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap = p_cap.add_run("Table 2.1: Comprehensive comparison of state-of-the-art RIS-assisted AirComp-FL frameworks.")
    r_cap.bold = True
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(10)

    add_heading_2(doc, "2.7 Research Gaps and Problem Formulation Rationale")
    p9 = (
        "A rigorous synthesis of the existing literature reveals three fundamental research gaps that motivate this project phase:\n\n"
        "1. Gap 1: Absence of Learning-Centric Active RIS Control. Existing active-RIS optimization schemes universally focus on instantaneous physical-layer "
        "metrics (such as SNR or estimated AirComp MSE). These physical metrics fail to capture how analog aggregation errors propagate through successive stochastic "
        "gradient updates across multiple federated learning rounds.\n\n"
        "2. Gap 2: Critical Blindness to Imperfect CSI. In current active-RIS AirComp literature, controllers treat estimated CSI as ground truth. Under imperfect CSI, "
        "this blind spot causes naive algorithms to drive active amplification to maximum gain, severely amplifying both surface thermal noise and channel estimation "
        "errors, leading to divergence in federated optimization.\n\n"
        "3. Gap 3: Excessive Computational Complexity of Existing Solvers. Standard robust beamforming approaches in multi-antenna RIS systems rely heavily on "
        "Semidefinite Relaxation (SDR) or Alternating Direction Method of Multipliers (ADMM). These matrix-lifting algorithms exhibit computational complexities "
        "scaling as O(N^4.5) to O(N^6), rendering them completely intractable for real-time edge execution within channel coherence time intervals.\n\n"
        "To bridge these critical gaps, this project formulates an analytical Convergence-Aware optimization framework that integrates CSI error statistics "
        "directly into the active RIS control loop and solves the joint problem with O(N^2) complexity, establishing an optimal balance between wireless aggregation "
        "quality and end-to-end federated learning convergence."
    )
    add_p(doc, p9)

    doc.add_page_break()
