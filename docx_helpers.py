import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Sets the background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets the internal padding of a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    """Applies clean borders to a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_header_footer(doc):
    """Configures running header and page numbering footer."""
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)  # Binding margin
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.27)   # A4 width
        section.page_height = Inches(11.69) # A4 height
        
        # Footer
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("Department of ECE, HICET  |  Page ")
        f_run.font.name = "Times New Roman"
        f_run.font.size = Pt(9)
        f_run.font.italic = True
        f_run.font.color.rgb = RGBColor(100, 100, 100)
        
        # Dynamic Page Field
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = "PAGE"
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'separate')
        fldChar3 = OxmlElement('w:fldChar')
        fldChar3.set(qn('w:fldCharType'), 'end')
        
        r = f_p.add_run()._r
        r.append(fldChar1)
        r.append(instrText)
        r.append(fldChar2)
        r.append(fldChar3)

def add_p(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=6, line_spacing=1.5, bold=False, italic=False, font_size=12, font_color=RGBColor(0,0,0)):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = "Times New Roman"
        run.font.size = Pt(font_size)
        run.font.color.rgb = font_color
    return p

def add_heading_1(doc, text, space_before=18, space_after=8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(15, 23, 42) # Dark navy
    return p

def add_heading_2(doc, text, space_before=14, space_after=4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(13.5)
    run.font.color.rgb = RGBColor(30, 41, 59) # Slate
    return p

def add_heading_3(doc, text, space_before=10, space_after=2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_callout(doc, title, text, bg_hex="F8FAFC", border_color="2563EB"):
    """Adds a highlighted callout box for theorems, definitions, and remarks."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    cell.width = Inches(6.0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    r_t = p.add_run(title + "\n")
    r_t.bold = True
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(11)
    r_t.font.color.rgb = RGBColor(15, 23, 42)
    
    r_b = p.add_run(text)
    r_b.italic = True
    r_b.font.name = "Times New Roman"
    r_b.font.size = Pt(10.5)
    r_b.font.color.rgb = RGBColor(30, 41, 59)
    
    # Empty space after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)

def add_equation(doc, eq_text, eq_number=""):
    """Adds an indented equation line with right-aligned numbering."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    
    r_eq = p.add_run(eq_text)
    r_eq.bold = False
    r_eq.italic = True
    r_eq.font.name = "Times New Roman"
    r_eq.font.size = Pt(11.5)
    
    if eq_number:
        r_num = p.add_run(f"\t\t\t\t({eq_number})")
        r_num.bold = True
        r_num.font.name = "Times New Roman"
        r_num.font.size = Pt(11)
        r_num.font.color.rgb = RGBColor(70, 70, 70)

def add_image_figure(doc, img_path, caption_num, caption_title, explanation="", width=Inches(5.4)):
    """Inserts a figure image with centered caption and explanatory text."""
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.paragraph_format.line_spacing = 1.15
        
        r_lbl = p_cap.add_run(f"Figure {caption_num}: ")
        r_lbl.bold = True
        r_lbl.font.name = "Times New Roman"
        r_lbl.font.size = Pt(10.5)
        r_lbl.font.color.rgb = RGBColor(15, 23, 42)
        
        r_title = p_cap.add_run(caption_title)
        r_title.bold = False
        r_title.font.name = "Times New Roman"
        r_title.font.size = Pt(10.5)
        
        if explanation:
            p_exp = doc.add_paragraph()
            p_exp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_exp.paragraph_format.space_before = Pt(2)
            p_exp.paragraph_format.space_after = Pt(8)
            p_exp.paragraph_format.line_spacing = 1.25
            r_exp = p_exp.add_run(explanation)
            r_exp.font.name = "Times New Roman"
            r_exp.font.size = Pt(10)
            r_exp.font.color.rgb = RGBColor(50, 50, 50)
    else:
        print(f"Warning: Image file not found: {img_path}")

print("Docx helper module ready.")
